"""Behavioral regressions for the shipped skill, without network or pip."""
import json
import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

SKILL = Path(__file__).resolve().parents[2] / "content/skills/open-kimi-ppt"
sys.path.insert(0, str(SKILL / "scripts"))
import runtime
from browser_export import allowed_request, ExportBrowser
from export_pptx import build_payload, patch_transitions, verify_output, PPTX_CONTENT_TYPE
import yaml
import websocket

HAS_OFFICE = importlib.util.find_spec("pptx") is not None


class ExportTests(unittest.TestCase):
    def test_dependencies_are_bundled(self):
        for module in (yaml, websocket):
            self.assertTrue(Path(module.__file__).is_relative_to(SKILL / "scripts/vendor"))
        self.assertEqual(yaml.safe_load('title: "中文"\npages: [pages/1.page]')["title"], "中文")
        from websocket._abnf import ABNF
        frame = ABNF.create_frame("中文", ABNF.OPCODE_TEXT)
        self.assertGreater(len(frame.format()), 6)

    def test_missing_components_have_actionable_errors(self):
        with patch("runtime.sys.platform", "linux"), patch("runtime.shutil.which", return_value=None):
            with self.assertRaisesRegex(runtime.ExportError, "BROWSER_UNAVAILABLE"):
                runtime.chrome_binary()
        with patch("runtime.importlib.util.find_spec", return_value=None), patch("runtime.office_candidates", return_value=[]):
            with self.assertRaisesRegex(runtime.ExportError, "OFFICE_RUNTIME_REQUIRED"):
                runtime.require_pillow()

    def test_external_uploads_and_telemetry_are_blocked(self):
        local = "http://127.0.0.1:12345"
        for url in ("https://www.kimi.com/apiv2/upload", "https://example.org/image.png",
                    "https://apmplus.volces.com/settings/get/webpro", "https://statics.kimi.ai/collect?data=secret"):
            for method in ("POST", "PUT", "GET", "OPTIONS"):
                self.assertFalse(allowed_request(url, method, local))
        self.assertFalse(allowed_request("https://statics.kimi.ai/neo-design/assets/a.js?payload=x", "GET", local))
        self.assertTrue(allowed_request("https://statics.kimi.ai/neo-design/assets/Editor-CFJyt8gP.js", "GET", local))
        self.assertFalse(allowed_request("https://statics.kimi.ai/neo-design/assets/private-document.js", "GET", local))
        self.assertTrue(allowed_request("https://statics.kimi.ai/neo-design-static/p-font/MiSans.woff2", "GET", local))
        self.assertTrue(allowed_request(local + "/payload.json", "GET", local))
        self.assertFalse(allowed_request(local + "/chrome-profile/Cookies", "GET", local))

    def test_payload_preserves_pages_and_rejects_paths_outside_project(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest = root / "deck.pptd"
            manifest.write_text('version: v2\ntitle: 测试\npages: [page.page]\n')
            (root / "page.page").write_text('elements: []\n')
            payload = build_payload(manifest)
            self.assertEqual(payload["title"], "测试")
            self.assertEqual(payload["pages"], [{"path": "page.page", "content": "elements: []\n"}])
            manifest.write_text('version: v2\npages: [../secret.page]\n')
            with self.assertRaisesRegex(runtime.ExportError, "escapes"):
                build_payload(manifest)
            manifest.write_text('version: v2\npages: [page.page]\n')
            (root / "leaked.png").symlink_to(Path(tmp).parent / "secret.png")
            # A readable image symlink outside the project must never enter the payload.
            with tempfile.TemporaryDirectory() as other:
                external = Path(other) / "secret.png"
                external.write_bytes(b"private")
                (root / "leaked.png").unlink()
                (root / "leaked.png").symlink_to(external)
                with self.assertRaisesRegex(runtime.ExportError, "escapes"):
                    build_payload(manifest)

    def test_browser_requests_drop_document_headers_and_cookies(self):
        browser = ExportBrowser("/tmp/test", "/tmp/test/downloads", "http://127.0.0.1:123/export_host.html")
        request = {"url": "https://statics.kimi.ai/neo-design/assets/Editor-CFJyt8gP.js", "method": "GET",
                   "headers": {"Cookie": "secret", "X-Document": "private", "Accept": "private",
                               "Origin": "https://www.kimi.com"}}
        with patch.object(browser, "call") as call:
            browser._event({"method": "Fetch.requestPaused", "sessionId": "iframe", "params": {"requestId": "req", "request": request}})
            headers = call.call_args.args[1]["headers"]
            self.assertNotIn("secret", json.dumps(headers))
            self.assertNotIn("private", json.dumps(headers))
            self.assertIn({"name": "Origin", "value": "https://www.kimi.com"}, headers)

    def test_pptx_transition_and_archive_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            pptx = Path(tmp) / "deck.pptx"
            with zipfile.ZipFile(pptx, "w") as archive:
                archive.writestr("[Content_Types].xml", PPTX_CONTENT_TYPE)
                archive.writestr("ppt/presentation.xml", "<presentation/>")
                archive.writestr("ppt/slides/slide1.xml", '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cSld><p:sp><p:extLst/></p:sp></p:cSld><p:timing/></p:sld>')
            self.assertEqual(patch_transitions(pptx, "fade"), 1)
            result = verify_output(pptx, "fade", False)
            self.assertEqual(result["fadeTransitions"], 1)
            patch_transitions(pptx, "none")
            self.assertEqual(verify_output(pptx, "none", False)["fadeTransitions"], 0)

    def test_idle_browser_wait_is_bounded(self):
        browser = ExportBrowser("/tmp/test", "/tmp/test/downloads", "http://127.0.0.1:123/export_host.html")
        with patch.object(browser, "pump"), patch.object(browser, "evaluate", return_value=None):
            with self.assertRaisesRegex(runtime.ExportError, "EDITOR_TIMEOUT"):
                browser.wait("false", "iframe", timeout=0.01)

    @unittest.skipUnless(HAS_OFFICE, "requires the managed office runtime")
    def test_real_local_export_keeps_native_editable_elements_without_network(self):
        from PIL import Image
        from pptx import Presentation
        from export_pptx import export_pptx
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            Image.new("RGB", (80, 40), "red").save(root / "image.png")
            manifest = root / "deck.pptd"
            manifest.write_text(yaml.safe_dump({"version": "v2", "title": "中文验证", "size": [960, 540],
                "theme": {"colors": {"primary": "#224466"}, "textStyles": {"title": {"fontSize": 32, "color": "$primary"}}},
                "pages": ["one.page"]}, allow_unicode=True))
            elements = [
                {"elementId": "gradient", "elementType": "shape", "bounds": [10, 10, 100, 50], "shapeName": "roundRect",
                 "fill": {"type": "gradient", "gradientType": "linear", "stops": [{"position": 0, "color": "#224466"}, {"position": 1, "color": "#4488AA"}]},
                 "border": {"color": "#000000", "width": 1}},
                {"elementId": "rich", "elementType": "text", "bounds": [120, 10, 400, 90],
                 "content": {"style": "$title", "fontFamily": {"latin": "Arial", "ea": "Microsoft YaHei"}, "text": '<p><span style="font-size:30px">销售</span><strong>增长</strong></p>'}},
                {"elementId": "image", "elementType": "image", "bounds": [10, 100, 100, 100], "src": "image.png", "fit": {"mode": "contain"}},
                {"elementId": "line", "elementType": "line", "bounds": [10, 220, 300, 1], "points": "0,0 1,1", "viewBox": [1, 1], "border": {"color": "#000000"}},
                {"elementId": "table", "elementType": "table", "bounds": [10, 250, 300, 200], "columnWidths": [0.5, 0.5], "rowHeights": [0.5, 0.5],
                 "rows": [[{"text": "合并", "colSpan": 2, "fill": {"type": "solid", "color": "#DDEEFF"}}], [{"text": "A"}, {"text": "B"}]]},
                {"elementId": "chart", "elementType": "chart", "bounds": [400, 200, 500, 300], "data": {"cols": ["类目", "值"], "rows": [["甲", 2], ["乙", 3]]},
                 "series": [{"type": "pie", "encode": {"category": "类目", "value": "值"}}], "title": "比例", "legend": True},
            ]
            (root / "one.page").write_text(yaml.safe_dump({"background": {"type": "solid", "color": "#FFFFFF"}, "elements": elements}, allow_unicode=True))
            with patch("socket.socket", side_effect=AssertionError("PPTX conversion must not use network")), \
                 patch("subprocess.Popen", side_effect=AssertionError("PPTX conversion must not launch a browser")):
                summary = export_pptx(manifest, root / "out.pptx", "fade", False)
            self.assertFalse(summary["networkUsed"])
            self.assertEqual(summary["fadeTransitions"], 1)
            prs = Presentation(str(root / "out.pptx"))
            shapes = {shape.name: shape for shape in prs.slides[0].shapes}
            self.assertEqual(len(shapes), 6)
            self.assertEqual(shapes["rich"].text, "销售增长")
            self.assertTrue(shapes["rich"].text_frame.paragraphs[0].runs[1].font.bold)
            self.assertTrue(shapes["table"].has_table)
            self.assertEqual(shapes["table"].table.cell(0, 0).text, "合并")
            self.assertTrue(shapes["chart"].has_chart)
            self.assertEqual(tuple(shapes["chart"].chart.series[0].values), (2.0, 3.0))
            self.assertAlmostEqual(shapes["image"].width / shapes["image"].height, 2)

    @unittest.skipUnless(HAS_OFFICE, "requires the managed office runtime")
    def test_unsupported_elements_do_not_replace_existing_output(self):
        from export_pptx import export_pptx
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "deck.pptd").write_text("version: v2\nsize: [960, 540]\npages: [one.page]\n")
            (root / "one.page").write_text("elements:\n- elementId: icon\n  elementType: icon\n  iconName: fas:house\n  bounds: [10, 10, 40, 40]\n")
            output = root / "out.pptx"
            output.write_bytes(b"old output")
            with self.assertRaisesRegex(runtime.ExportError, "UNSUPPORTED_PPTD_FEATURE"):
                export_pptx(root / "deck.pptd", output, "fade", False, force=True)
            self.assertEqual(output.read_bytes(), b"old output")
            self.assertEqual(list(root.glob("*.tmp")), [])


if __name__ == "__main__":
    unittest.main()
