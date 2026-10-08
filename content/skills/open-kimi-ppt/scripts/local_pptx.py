"""Offline PPTD v2 -> editable PowerPoint using EvoWork's office runtime.

Supports text, preset shapes, local raster images, straight lines, native tables
and homogeneous category charts. Unsupported features fail with their element
ID instead of silently dropping or rasterizing editable content.
"""
import base64
from html.parser import HTMLParser
from io import BytesIO
import math
from pathlib import Path
import re

from runtime import ExportError
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt


def unsupported(feature):
    raise ExportError(f"UNSUPPORTED_PPTD_FEATURE：本地导出暂不支持 {feature}；请按 reference/evowork-local-export.md 调整，完整 PPTD 保留。")


def number(value, name, positive=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or (positive and value <= 0):
        raise ExportError(f"INVALID_PPTD：{name} 必须是{'正' if positive else '有限'}数值。")
    return value


def color(value, theme):
    seen = set()
    while isinstance(value, str) and value.startswith("$"):
        if value in seen:
            raise ExportError("INVALID_PPTD：主题颜色循环引用。")
        seen.add(value)
        value = theme.get("colors", {}).get(value[1:])
    if not isinstance(value, str) or not re.fullmatch(r"#[0-9a-fA-F]{6}(?:[0-9a-fA-F]{2})?", value):
        raise ExportError(f"INVALID_PPTD：颜色必须是 #RRGGBB 或 #RRGGBBAA：{value}")
    return value[1:7], int(value[7:9], 16) / 255 if len(value) == 9 else 1


def alpha(node, opacity):
    if opacity < 1:
        item = OxmlElement("a:alpha")
        item.set("val", str(round(opacity * 100000)))
        node.append(item)


def set_fill(fill, spec, theme, opacity=1):
    if not spec:
        fill.background()
    elif spec.get("type") == "solid":
        fill.solid()
        rgb, transparency = color(spec["color"], theme)
        fill.fore_color.rgb = RGBColor.from_string(rgb)
        alpha(fill.fore_color._color._xClr, opacity * transparency)
    elif spec.get("type") == "gradient" and spec.get("gradientType", "linear") == "linear":
        stops = spec.get("stops", [])
        if len(stops) < 2:
            raise ExportError("INVALID_PPTD：渐变至少需要两个颜色节点。")
        fill.gradient()
        grad = fill._fill._gradFill
        stop_list = grad.gsLst
        for child in list(stop_list):
            stop_list.remove(child)
        for stop in stops:
            pos = stop.get("offset", stop.get("position"))
            if not isinstance(pos, (float, int)) or not 0 <= pos <= 1:
                raise ExportError("INVALID_PPTD：渐变节点位置应为 0 到 1。")
            item = OxmlElement("a:gs")
            item.set("pos", str(round(pos * 100000)))
            rgb, transparency = color(stop["color"], theme)
            item_color = OxmlElement("a:srgbClr")
            item_color.set("val", rgb)
            alpha(item_color, opacity * transparency)
            item.append(item_color)
            stop_list.append(item)
        fill.gradient_angle = (360 - spec.get("angle", 0)) % 360
    else:
        unsupported("这种填充方式（图片背景请使用独立 image 元素）")


def set_line(shape, border, theme, opacity=1):
    if not border:
        shape.line.fill.background()
        return
    if border.get("style", "solid") not in ("solid", "dash", "dot"):
        unsupported("边框线型 " + str(border.get("style")))
    rgb, transparency = color(border.get("color", "#000000"), theme)
    shape.line.color.rgb = RGBColor.from_string(rgb)
    shape.line.width = Pt(number(border.get("width", 1), "border.width"))
    alpha(shape.line.color._color._xClr, opacity * transparency)
    if border.get("style") in ("dash", "dot"):
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH if border["style"] == "dash" else MSO_LINE_DASH_STYLE.ROUND_DOT


CSS_KEYS = {"font-size": "fontSize", "font-family": "fontFamily", "color": "color",
            "text-align": "textAlign", "line-height": "lineHeight", "letter-spacing": "letterSpacing",
            "margin-top": "marginTop", "background-color": "backgroundColor"}


def css_style(value):
    result = {}
    for declaration in value.split(";"):
        key, sep, item = declaration.partition(":")
        if not sep:
            continue
        key, item = key.strip(), item.strip()
        if key in CSS_KEYS:
            target = CSS_KEYS[key]
            if key in ("font-size", "letter-spacing", "margin-top", "line-height"):
                if item.endswith("px") or item.endswith("pt"):
                    if key == "line-height":
                        target = "lineHeightPx"
                    item = item[:-2]
                result[target] = float(item)
            else:
                result[target] = item.strip("\"'")
        elif key == "font-weight":
            result["bold"] = item in ("bold", "bolder") or (item.isdigit() and int(item) >= 600)
        elif key == "font-style":
            result["italic"] = item == "italic"
        elif key == "text-decoration":
            result["underline"] = "underline" in item
            result["strike"] = "line-through" in item
        else:
            unsupported("文本 CSS 属性 " + key)
    return result


class RichText(HTMLParser):
    def __init__(self, text_frame, style, theme, opacity=1):
        super().__init__(convert_charrefs=True)
        text_frame.clear()
        self.frame = text_frame
        self.style = style
        self.theme = theme
        self.opacity = opacity
        self.stack = []
        self.paragraph = text_frame.paragraphs[0]
        self.started = False
        self.new_paragraph = False

    def current_style(self):
        result = dict(self.style)
        for _, item in self.stack:
            result.update(item)
        return result

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "br":
            self.paragraph.add_line_break()
            return
        if tag not in ("p", "div", "span", "strong", "b", "em", "i", "u", "s", "strike", "sup", "sub", "ul", "ol", "li", "a"):
            unsupported("富文本标签 " + tag)
        if tag in ("p", "div", "li"):
            if self.started or self.new_paragraph:
                self.paragraph = self.frame.add_paragraph()
            self.started = True
            self.new_paragraph = False
        style = css_style(attrs.get("style", ""))
        for tags, key in ((('strong', 'b'), 'bold'), (('em', 'i'), 'italic'), (('u',), 'underline'), (('s', 'strike'), 'strike')):
            if tag in tags:
                style[key] = True
        if tag in ("sup", "sub"):
            style["baseline"] = 30000 if tag == "sup" else -25000
        if tag == "a" and attrs.get("href"):
            style["hyperlink"] = attrs["href"]
        self.stack.append((tag, style))
        if tag == "li":
            bullet = OxmlElement("a:buChar")
            bullet.set("char", "•")
            self.paragraph._p.get_or_add_pPr().append(bullet)

    def handle_endtag(self, tag):
        for index in range(len(self.stack) - 1, -1, -1):
            if self.stack[index][0] == tag:
                self.stack = self.stack[:index]
                break
        if tag in ("p", "div", "li"):
            self.new_paragraph = True

    def handle_data(self, text):
        if not text or (not text.strip() and self.new_paragraph):
            return
        style = self.current_style()
        paragraph = self.paragraph
        paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
                               "justify": PP_ALIGN.JUSTIFY}.get(style.get("textAlign", style.get("align", ["left"])[0]), PP_ALIGN.LEFT)
        if "lineHeightPx" in style:
            paragraph.line_spacing = Pt(style["lineHeightPx"])
        elif "lineHeight" in style:
            paragraph.line_spacing = float(style["lineHeight"])
        paragraph.space_before = Pt(style.get("marginTop", 0))
        run = paragraph.add_run()
        run.text = text
        font = run.font
        family = style.get("fontFamily", "Arial")
        font.name = family.get("latin", "Arial") if isinstance(family, dict) else family
        font.size = Pt(number(style.get("fontSize", 18), "fontSize", True))
        font.bold = bool(style.get("bold", False))
        font.italic = bool(style.get("italic", False))
        font.underline = bool(style.get("underline", False))
        rpr = run._r.get_or_add_rPr()
        if style.get("strike"):
            rpr.set("strike", "sngStrike")
        if "baseline" in style:
            rpr.set("baseline", str(style["baseline"]))
        if "letterSpacing" in style:
            rpr.set("spc", str(round(style["letterSpacing"] * 100)))
        # Set East Asian typeface too, so Chinese text does not fall back to a Latin-only face.
        ea = OxmlElement("a:ea")
        ea.set("typeface", family.get("ea", font.name) if isinstance(family, dict) else font.name)
        rpr.append(ea)
        rgb, transparency = color(style.get("color", "#000000"), self.theme)
        font.color.rgb = RGBColor.from_string(rgb)
        alpha(font.color._color._xClr, transparency * self.opacity)
        if style.get("backgroundColor"):
            rgb, _ = color(style["backgroundColor"], self.theme)
            highlight = OxmlElement("a:highlight")
            clr = OxmlElement("a:srgbClr")
            clr.set("val", rgb)
            highlight.append(clr)
            rpr.append(highlight)
        if style.get("hyperlink"):
            run.hyperlink.address = style["hyperlink"]
        self.started = True


def render_text(frame, content, theme, opacity=1):
    style = {**theme.get("textStyles", {}).get(str(content.get("style", "")).lstrip("$"), {}), **content}
    for key in ("gradient", "shadow"):
        if style.get(key):
            unsupported("文本 " + key)
    frame.margin_top = frame.margin_bottom = frame.margin_left = frame.margin_right = 0
    frame.word_wrap = style.get("wrap", True)
    frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(style.get("align", ["left", "top"])[-1], MSO_ANCHOR.TOP)
    if style.get("textDirection") == "vertical":
        frame._txBody.bodyPr.set("vert", "eaVert")
    parser = RichText(frame, style, theme, opacity)
    parser.feed(str(style.get("text", "")))
    parser.close()


def add_image(slide, element, root):
    from export_pptx import safe_project_path
    src = element.get("src", "")
    if src.startswith("data:image/"):
        image = Image.open(BytesIO(base64.b64decode(src.split(",", 1)[1], validate=True)))
    else:
        if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", src):
            unsupported("远程图片（请先下载到 media/）")
        image = Image.open(safe_project_path(root, src))
    image.load()
    crop = element.get("crop", {})
    width, height = image.size
    image = image.crop((round(width * crop.get("left", 0)), round(height * crop.get("top", 0)),
                        round(width * (1 - crop.get("right", 0))), round(height * (1 - crop.get("bottom", 0)))))
    x, y, w, h = element["bounds"]
    mode = element.get("fit", {}).get("mode", "cover")
    if mode == "cover":
        ratio = max(w / image.width, h / image.height)
        visible_w, visible_h = w / ratio, h / ratio
        left, top = (image.width - visible_w) / 2, (image.height - visible_h) / 2
        image = image.crop((round(left), round(top), round(left + visible_w), round(top + visible_h)))
    elif mode == "contain":
        ratio = min(w / image.width, h / image.height)
        pw, ph = image.width * ratio, image.height * ratio
        x, y, w, h = x + (w - pw) / 2, y + (h - ph) / 2, pw, ph
    elif mode != "fill":
        unsupported("图片 fit.mode=" + str(mode))
    if element.get("opacity", 1) != 1:
        image = image.convert("RGBA")
        image.putalpha(image.getchannel("A").point(lambda v: round(v * element["opacity"])))
    buffer = BytesIO()
    image.save(buffer, "PNG")
    buffer.seek(0)
    shape = slide.shapes.add_picture(buffer, *[Pt(v) for v in (x, y, w, h)])
    crop_shape = element.get("cropShape", {})
    if crop_shape:
        shape.auto_shape_type = MSO_SHAPE.from_xml(crop_shape.get("shapeName", "rect"))
    return shape


def add_table(slide, element, theme):
    rows = element["rows"]
    col_widths, row_heights = element["columnWidths"], element["rowHeights"]
    x, y, width, height = element["bounds"]
    if len(rows) != len(row_heights) or not rows or not col_widths:
        raise ExportError("INVALID_PPTD：表格行列数量错误。")
    for ratios in (col_widths, row_heights):
        if not all(isinstance(v, (int, float)) and v > 0 for v in ratios) or abs(sum(ratios) - 1) > 0.001:
            raise ExportError("INVALID_PPTD：表格行列比例必须为正且合计为 1。")
    shape = slide.shapes.add_table(len(rows), len(col_widths), *[Pt(v) for v in element["bounds"]])
    table = shape.table
    for i, ratio in enumerate(col_widths):
        table.columns[i].width = Pt(width * ratio)
    for i, ratio in enumerate(row_heights):
        table.rows[i].height = Pt(height * ratio)
    raw_style = element.get("style", {})
    style = theme.get("tableStyles", {}).get(raw_style.lstrip("$"), {}) if isinstance(raw_style, str) else raw_style
    occupied = set()
    for r, row in enumerate(rows):
        c = 0
        for source in row:
            while (r, c) in occupied:
                c += 1
            rs, cs = source.get("rowSpan", 1), source.get("colSpan", 1)
            if c + cs > len(col_widths) or r + rs > len(rows):
                raise ExportError("INVALID_PPTD：合并单元格超出表格。")
            cell = table.cell(r, c)
            if rs > 1 or cs > 1:
                cell.merge(table.cell(r + rs - 1, c + cs - 1))
            occupied.update((rr, cc) for rr in range(r, r + rs) for cc in range(c, c + cs))
            row_style = style.get("firstRowStyle", {}) if r == 0 else style.get("lastRowStyle", {}) if r == len(rows) - 1 else {}
            column_style = style.get("firstColumnStyle", {}) if c == 0 else style.get("lastColumnStyle", {}) if c == len(col_widths) - 1 else {}
            cycle = style.get("bodyStyles", [])
            baseline = {**style.get("cellStyle", {}), **(cycle[(r - 1) % len(cycle)] if cycle and 0 < r < len(rows) - 1 else {})}
            baseline.update({**column_style, **row_style} if style.get("rowOverColumn", True) else {**row_style, **column_style})
            baseline.update(theme.get("textStyles", {}).get(str(source.get("textStyle", "")).lstrip("$"), {}))
            merged = {"align": ["center", "middle"], **baseline, **source}
            set_fill(cell.fill, merged.get("fill", element.get("fill")), theme)
            render_text(cell.text_frame, merged, theme)
            # Explicit border is rendered on all four sides; null removes a side.
            border = merged.get("border", {"style": "solid", "width": 1, "color": "#000000"})
            sides = border if isinstance(border, list) else [border] * 4
            tcpr = cell._tc.get_or_add_tcPr()
            for name, spec in zip(("lnT", "lnR", "lnB", "lnL"), sides):
                line = OxmlElement("a:" + name)
                if spec:
                    line.set("w", str(Pt(spec.get("width", 1))))
                    solid = OxmlElement("a:solidFill")
                    clr = OxmlElement("a:srgbClr")
                    clr.set("val", color(spec.get("color", "#000000"), theme)[0])
                    solid.append(clr)
                    line.append(solid)
                else:
                    line.append(OxmlElement("a:noFill"))
                tcpr.append(line)
            c += cs
    return shape


def add_chart(slide, element, theme):
    if set(element) - {"elementId", "elementType", "bounds", "data", "series", "title", "legend"}:
        unsupported("高级图表配置（本地支持 data/series/title/legend）")
    series = element["series"]
    kinds = {s["type"] for s in series}
    kinds_map = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE,
                 "area": XL_CHART_TYPE.AREA, "pie": XL_CHART_TYPE.PIE, "radar": XL_CHART_TYPE.RADAR}
    if len(kinds) != 1 or not kinds.issubset(kinds_map):
        unsupported("混合图或图表类型 " + ",".join(sorted(kinds)))
    columns = element["data"]["cols"]
    rows = element["data"]["rows"]
    if not rows or len(columns) != len(set(columns)) or any(len(row) != len(columns) for row in rows):
        raise ExportError("INVALID_PPTD：图表数据列无效。")
    kind = next(iter(kinds))
    encoding = series[0].get("encode", {})
    category = encoding.get("x", encoding.get("category"))
    if not category or category not in columns:
        unsupported("缺少分类列的图表")
    data = CategoryChartData()
    data.categories = [str(row[columns.index(category)]) for row in rows]
    for item in series:
        if set(item) - {"type", "encode", "name"}:
            unsupported("高级图表系列配置")
        encoding = item.get("encode", {})
        if encoding.get("x", encoding.get("category")) != category:
            unsupported("使用不同分类列的多系列图表")
        value = encoding.get("y", encoding.get("value"))
        if value not in columns:
            raise ExportError("INVALID_PPTD：图表数值列不存在。")
        values = [row[columns.index(value)] for row in rows]
        data.add_series(item.get("name", value), [None if v is None else float(v) for v in values])
    shape = slide.shapes.add_chart(kinds_map[kind], *[Pt(v) for v in element["bounds"]], data)
    chart = shape.chart
    title = element.get("title")
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title if isinstance(title, str) else title.get("text", "")
    if not isinstance(element.get("legend", True), bool):
        unsupported("图例样式配置（请使用 true/false）")
    chart.has_legend = bool(element.get("legend", kind == "pie"))
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    return shape


def write_pptx(manifest, output):
    from export_pptx import read_yaml_mapping, safe_project_path
    _, document = read_yaml_mapping(manifest)
    size = document.get("size", [960, 540])
    if len(size) != 2:
        raise ExportError("INVALID_PPTD：页面尺寸必须包含宽高。")
    prs = Presentation()
    prs.slide_width, prs.slide_height = [Pt(number(v, "size", True)) for v in size]
    prs.core_properties.title = str(document.get("title", manifest.stem))
    theme = document.get("theme", {})
    if document.get("customFonts"):
        unsupported("customFonts 远程字体配置（请使用客户已有的系统字体）")
    root = manifest.parent.resolve()
    for filename in document["pages"]:
        _, page = read_yaml_mapping(safe_project_path(root, filename))
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if page.get("animations") or page.get("animation"):
            unsupported("页内动画")
        background = page.get("background")
        if background and background.get("type") == "image":
            add_image(slide, {**background, "bounds": [0, 0, *size]}, root)
        elif background:
            set_fill(slide.background.fill, background, theme)
        seen = set()
        for element in page["elements"]:
            ident = element.get("elementId", "")
            if not ident or ident in seen:
                raise ExportError("INVALID_PPTD：元素 ID 缺失或重复。")
            seen.add(ident)
            try:
                bounds = element["bounds"]
                if len(bounds) != 4:
                    raise ExportError("INVALID_PPTD：bounds 必须包含 x,y,width,height。")
                for index, v in enumerate(bounds):
                    number(v, "bounds", index >= 2)
                for field in ("shadow", "animation", "animations"):
                    if element.get(field):
                        unsupported(field)
                kind = element["elementType"]
                opacity = element.get("opacity", 1)
                if not 0 <= opacity <= 1:
                    raise ExportError("INVALID_PPTD：opacity 应为 0 到 1。")
                if kind == "text":
                    shape = slide.shapes.add_textbox(*[Pt(v) for v in bounds])
                    render_text(shape.text_frame, element["content"], theme, opacity)
                elif kind == "shape":
                    try:
                        preset = MSO_SHAPE.from_xml(element["shapeName"])
                    except ValueError:
                        unsupported("形状 " + element["shapeName"])
                    shape = slide.shapes.add_shape(preset, *[Pt(v) for v in bounds])
                    for i, adjustment in enumerate(element.get("adjustments", [])):
                        shape.adjustments[i] = adjustment / 100000
                    set_fill(shape.fill, element.get("fill"), theme, opacity)
                elif kind == "image":
                    shape = add_image(slide, element, root)
                elif kind == "line":
                    points = [tuple(map(float, pair.split(","))) for pair in element["points"].split()]
                    if len(points) != 2:
                        unsupported("曲线（请使用两点直线）")
                    vx, vy = element["viewBox"]
                    x, y, w, h = bounds
                    x1, y1 = points[0]
                    x2, y2 = points[1]
                    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x + x1 / vx * w), Pt(y + y1 / vy * h), Pt(x + x2 / vx * w), Pt(y + y2 / vy * h))
                    for name, arrow in zip(("headEnd", "tailEnd"), element.get("arrow", [None, None])):
                        if arrow:
                            node = OxmlElement("a:" + name)
                            node.set("type", "triangle" if arrow == "arrow" else arrow)
                            shape.line._get_or_add_ln().append(node)
                elif kind == "table":
                    shape = add_table(slide, element, theme)
                elif kind == "chart":
                    shape = add_chart(slide, element, theme)
                else:
                    unsupported("元素类型 " + str(kind))
                shape.name = ident
                if kind not in ("table", "chart"):
                    shape.rotation = element.get("rotation", 0)
                    if any(element.get("flip", [])):
                        xfrm = shape._element.spPr.xfrm
                        for name, value in zip(("flipH", "flipV"), element["flip"]):
                            xfrm.set(name, "1" if value else "0")
                    set_line(shape, element.get("border"), theme, opacity)
            except ExportError as error:
                raise ExportError(f"{filename} / {ident}: {error}") from error
            except (KeyError, ValueError, TypeError, IndexError, ZeroDivisionError) as error:
                raise ExportError(f"INVALID_PPTD：{filename} / {ident} ({type(error).__name__})") from error
    prs.save(str(output))
